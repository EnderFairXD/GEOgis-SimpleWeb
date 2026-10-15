"""Importación de los ficheros oficiales de calidad del aire del MITECO.

Fuente: Ministerio para la Transición Ecológica y el Reto Demográfico (MITECO),
«Datos oficiales de calidad del aire». Reutilización según la Ley 37/2007.
"""
import csv
import datetime
import io
import shutil
import unicodedata
import urllib.parse
import urllib.request
import zipfile
from collections import defaultdict
from pathlib import Path

from django.contrib.gis.geos import Point
from django.db import transaction
from openpyxl import load_workbook

from .models import Contaminante, Estacion, MedicionMensual

URL_MITECO_2024 = ('https://www.miteco.gob.es/content/dam/miteco/es/calidad-y-evaluacion-ambiental/'
                   'sgalsi/atm%c3%b3sfera-y-calidad-del-aire/evaluaci%c3%b3n-2024/')
URL_METAINFORMACION = URL_MITECO_2024 + 'Metainformacion_2024.xlsx'
URLS_MEDICIONES = {
    2024: (URL_MITECO_2024 + 'Datos_HH_2024.zip', URL_MITECO_2024 + 'Datos_DD_2024.zip'),
}


def descargar(url, carpeta):
    """Descarga `url` en `carpeta` y devuelve la ruta del fichero."""
    destino = Path(carpeta) / Path(urllib.parse.unquote(url)).name
    with urllib.request.urlopen(url, timeout=120) as respuesta, open(destino, 'wb') as fichero:
        shutil.copyfileobj(respuesta, fichero)
    return destino


def _sin_tildes(texto):
    normalizado = unicodedata.normalize('NFKD', str(texto).strip().upper())
    return ''.join(c for c in normalizado if not unicodedata.combining(c))


def _fecha(valor):
    if valor in (None, ''):
        return None
    if isinstance(valor, datetime.datetime):
        return valor.date()
    return datetime.date.fromisoformat(str(valor)[:10])


def _filas(hoja):
    """Filas de una hoja de Excel como diccionarios, usando la primera fila como cabecera."""
    filas = hoja.iter_rows(values_only=True)
    cabecera = next(filas)
    for fila in filas:
        if fila and fila[0] is not None:
            yield dict(zip(cabecera, fila))


@transaction.atomic
def importar_estaciones(ruta):
    """Carga el catálogo de contaminantes y las estaciones.

    Las estaciones con la ubicación corregida a mano conservan su posición; el resto de sus datos
    sí se actualiza. Devuelve (contaminantes, estaciones, ubicaciones corregidas conservadas).
    """
    libro = load_workbook(ruta, read_only=True)
    hoja_estaciones = next(h for h in libro.worksheets if h.title.startswith('Estaciones'))

    contaminantes = 0
    for fila in _filas(libro['Magnitudes']):
        Contaminante.objects.update_or_create(codigo=fila['MAGNITUD'], defaults={
            'abreviatura': fila['ABREVIATURA'],
            'nombre': fila['NOMBRE'],
            'unidad': fila['UNIDAD'],
        })
        contaminantes += 1

    corregidas = set(Estacion.objects.filter(ubicacion_corregida=True).values_list('codigo', flat=True))
    estaciones = conservadas = 0
    for fila in _filas(hoja_estaciones):
        datos = {
            'codigo_europeo': fila['COD_ESTACION_DEM'],
            'nombre': fila['NOMBRE'],
            'direccion': fila['DIRECCION'] or '',
            'municipio': fila['N_MUNICIPIO'],
            'provincia': fila['N_PROVINCIA'],
            'comunidad': fila['N_CCAA'],
            'red': fila['N_RED'],
            'tipo': _sin_tildes(fila['TIPO_ESTACION']),
            'area': _sin_tildes(fila['TIPO_AREA']),
            'altitud': fila['ALTITUD'],
            'fecha_alta': _fecha(fila['FECHA_INI']),
            'fecha_baja': _fecha(fila['FECHA_FIN']),
            'ubicacion': Point(float(fila['LONGITUD_G']), float(fila['LATITUD_G']), srid=4326),
        }
        if fila['COD_LOCAL'] in corregidas:
            del datos['ubicacion']
            conservadas += 1
        Estacion.objects.update_or_create(codigo=fila['COD_LOCAL'], defaults=datos)
        estaciones += 1

    return contaminantes, estaciones, conservadas


def _sumar_zip(ruta, prefijo):
    """Suma y cuenta los valores de cada estación, magnitud y mes de un ZIP del MITECO.

    `prefijo` es la letra de las columnas de valores: 'H' (horas, H01-H24) o 'D' (días, D01-D31).
    Se descartan los valores negativos (ruido de los analizadores). Si una estación tiene varios
    puntos de muestreo para la misma magnitud, se promedian juntos.
    """
    acumulado = defaultdict(lambda: [0.0, 0])
    with zipfile.ZipFile(ruta) as zf:
        for nombre in zf.namelist():
            with zf.open(nombre) as binario:
                lector = csv.DictReader(io.TextIOWrapper(binario, encoding='utf-8-sig'), delimiter=';')
                columnas = [c for c in lector.fieldnames if c[0] == prefijo and c[1:].isdigit()]
                for fila in lector:
                    codigo = int(fila['PROVINCIA']) * 1_000_000 + int(fila['MUNICIPIO']) * 1000 + int(fila['ESTACION'])
                    clave = (codigo, int(fila['MAGNITUD']), int(fila['ANNO']), int(fila['MES']))
                    for columna in columnas:
                        valor = fila[columna]
                        if valor and float(valor) >= 0:
                            acumulado[clave][0] += float(valor)
                            acumulado[clave][1] += 1
    return acumulado


@transaction.atomic
def importar_mediciones(ruta_horario, ruta_diario):
    """Calcula y guarda las medias mensuales. Los datos horarios tienen prioridad: los diarios
    solo se usan para los meses sin datos horarios (metales, benzo(a)pireno, captadores manuales).

    Devuelve un diccionario con lo importado y lo descartado.
    """
    sumas = _sumar_zip(ruta_horario, 'H')
    for clave, suma in _sumar_zip(ruta_diario, 'D').items():
        sumas.setdefault(clave, suma)

    estaciones = dict(Estacion.objects.values_list('codigo', 'pk'))
    contaminantes = dict(Contaminante.objects.values_list('codigo', 'pk'))
    estaciones_desconocidas, magnitudes_desconocidas = set(), set()
    mediciones, medidos = [], set()

    for (codigo, magnitud, anno, mes), (total, cuantos) in sumas.items():
        if codigo not in estaciones:
            estaciones_desconocidas.add(codigo)
            continue
        if magnitud not in contaminantes:
            magnitudes_desconocidas.add(magnitud)
            continue
        estacion_id, contaminante_id = estaciones[codigo], contaminantes[magnitud]
        mediciones.append(MedicionMensual(
            estacion_id=estacion_id, contaminante_id=contaminante_id,
            anno=anno, mes=mes, media=total / cuantos, num_datos=cuantos,
        ))
        medidos.add((estacion_id, contaminante_id))

    MedicionMensual.objects.bulk_create(
        mediciones, batch_size=5000, update_conflicts=True,
        unique_fields=['estacion', 'contaminante', 'anno', 'mes'], update_fields=['media', 'num_datos'],
    )
    Mide = Estacion.contaminantes.through
    Mide.objects.bulk_create(
        [Mide(estacion_id=e, contaminante_id=c) for e, c in medidos], ignore_conflicts=True,
    )

    return {
        'mediciones': len(mediciones),
        'estaciones': len({e for e, _ in medidos}),
        'estaciones_desconocidas': len(estaciones_desconocidas),
        'magnitudes_desconocidas': len(magnitudes_desconocidas),
    }
